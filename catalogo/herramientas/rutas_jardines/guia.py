"""Listas por jardín (Excel) y guía de campo tipo catálogo (un solo HTML con fotos embebidas).

Especie de cada árbol: binomio del nombre científico 2018; si falta, el binomio que más se usa
para ese nombre común 2014 en todo el inventario (si domina con 80% o más). Para emparejar con
nuestras fotos se usa el taxón de iNaturalist cuando existe (une sinónimos), si no el binomio.

Uso (desde esta carpeta):  python guia.py
Salida (gitignorada): catalogo/salida/inventario_arboles/Listas_por_jardin.xlsx y
Guia_arboles_por_jardin.html
"""
import base64
import csv
import io
import json
import re
import sqlite3
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageOps

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI.parents[1]))  # catalogo/, para importar comun/
from comun import rutas  # noqa: E402
from comun.rutas import RUTA_BD, RUTA_CONFIG  # noqa: E402
SALIDA = rutas.SALIDA / "inventario_arboles"
CSV_ARCGIS = SALIDA / "arcgis_2014_2018_export.csv"
XLSX_MAESTRO = SALIDA / "Inventario_Maestro.xlsx"
FICHAS = rutas.RUTA_FICHAS
PLANTILLA = AQUI / "plantilla_guia.html"
XLSX_SALIDA = SALIDA / "Listas_por_jardin.xlsx"
HTML_SALIDA = SALIDA / "Guia_arboles_por_jardin.html"

PARTES = ["arbol_completo", "tronco", "hoja", "flor", "fruto"]  # como en partes_arbol_etiquetas
LADO_MINIATURA = 360
RASGOS = ["altura", "tronco", "copa", "ramas", "hojas", "flores", "frutos"]


def norm(s):
    return unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower().strip()


# Erratas del inventario que iNaturalist no resuelve solo
ERRATAS = {"Eucaliptus citriodora": "Eucalyptus citriodora", "Melia azederach": "Melia azedarach",
           "Hibiscus rosasinensis": "Hibiscus rosa-sinensis"}


def binomio(s):
    w = (s or "").replace("×", " ").replace(";", " ").split()
    b = f"{w[0].capitalize()} {w[1].lower()}" if len(w) >= 2 else ""
    return ERRATAS.get(b, b)


def jardin_limpio(j):
    return j.replace("CamellÛn", "Camellón").replace("Area", "Área").replace("Tecnologico", "Tecnológico")


def texto_limpio(s):
    return (s or "").strip().replace("·", "á").replace("Ò", "ñ").replace("?", "á") if s else ""


def leer_inventario():
    filas = [r for r in csv.DictReader(open(CSV_ARCGIS, encoding="utf-8-sig")) if r["ID_J"] != "0"]
    comun_a_bi = defaultdict(Counter)
    for r in filas:
        b = binomio(r["Nom_cien"])
        if b:
            comun_a_bi[norm(r["Especie"])][b] += 1
    dominante = {c: cnt.most_common(1)[0][0] for c, cnt in comun_a_bi.items()
                 if cnt.most_common(1)[0][1] / sum(cnt.values()) >= 0.8}
    arboles = []
    for r in filas:
        b = binomio(r["Nom_cien"])
        c14 = texto_limpio(r["Especie"])
        dom = dominante.get(norm(r["Especie"]))
        aviso = ""
        if b and dom and dom != b and norm(r["Especie"]) != "especie no asignada":
            aviso = f"En 2014 decía «{c14}», que en el resto del inventario es {dom}"
        especie = b or dom or ""
        arboles.append({
            "placa": r["Placa"], "jardin": jardin_limpio(r["Jardín"]), "comun_2014": c14,
            "comun_2018": texto_limpio(r["Nom_comun"]), "cientifico_2018": r["Nom_cien"].strip(),
            "especie": especie, "aviso": aviso, "anio": r["Año_plant"].strip(),
            "altura": r["Altura2018"].strip(), "diam": r["DmFuste201"].strip(),
            "estado": r["Estado2018"].strip(), "obs": r["Obs"].strip(),
            "lat": r["Latitud"], "lon": r["Longitud"],
        })
    return arboles


def leer_contemplativo():
    ws = openpyxl.load_workbook(XLSX_MAESTRO, read_only=True)["Jardín Contemplativo (2-oct)"]
    filas = ws.iter_rows(values_only=True)
    i = {k: n for n, k in enumerate(next(filas))}
    return [f for f in filas if f[i["latitud"]] is not None], i


def fotos_por_especie(llave_de):
    """Por llave de especie: la mejor foto (la más pesada) de cada parte, ya en miniatura base64."""
    espejo = json.loads(RUTA_CONFIG.read_text(encoding="utf-8")).get("ruta_espejo_local_fotos")
    con = sqlite3.connect(RUTA_BD)
    filas = con.execute("""
        SELECT f.id, f.ruta_actual, f.tamano_bytes, e.nombre_cientifico, p.parte
        FROM partes_arbol_etiquetas p JOIN fotos f ON f.id = p.foto_id JOIN especies e ON e.id = f.especie_id
        WHERE f.estado = 'usable' AND f.id NOT IN (SELECT foto_id FROM archivos_no_procesables)
    """).fetchall()
    total = Counter(llave_de(binomio(n)) for (n,) in con.execute(
        "SELECT e.nombre_cientifico FROM fotos f JOIN especies e ON e.id = f.especie_id WHERE f.estado = 'usable'"))
    mejor = {}
    for fid, ruta, peso, nombre, parte in filas:
        if parte not in PARTES:
            continue
        k = (llave_de(binomio(nombre)), parte)
        local = Path(espejo) / Path(ruta).name if espejo else None
        if local and local.exists() and (k not in mejor or (peso or 0) > mejor[k][1]):
            mejor[k] = (local, peso or 0, fid)
    fotos = defaultdict(dict)
    for (llave, parte), (ruta, _, fid) in mejor.items():
        try:
            im = ImageOps.exif_transpose(Image.open(ruta)).convert("RGB")
        except OSError:
            continue
        im.thumbnail((LADO_MINIATURA, LADO_MINIATURA))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=62, optimize=True)
        fotos[llave][parte] = {"src": "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode(),
                               "foto_id": fid}
    return fotos, total


def construir():
    arboles = leer_inventario()
    con = sqlite3.connect(RUTA_BD)
    # inat_nombres_resueltos guarda el binomio en minúsculas
    taxon = {b.lower(): t for b, t in con.execute("SELECT binomio, taxon_id FROM inat_nombres_resueltos WHERE taxon_id IS NOT NULL")}
    llave_de = lambda b: f"t{taxon[b.lower()]}" if b.lower() in taxon else b  # noqa: E731

    fichas = {}
    for v in json.loads(FICHAS.read_text(encoding="utf-8")).values():
        if v["tipo"] == "arbol":
            c = v["campos"]
            fichas[llave_de(binomio(c.get("nombre_cientifico", "")))] = {
                "url": v["url"], "familia": c.get("familia", ""), "comun": c.get("nombre_comun", ""),
                "origen": c.get("lugar_de_origen", ""), "rasgos": {k: c[k] for k in RASGOS if c.get(k)}}

    print("Haciendo miniaturas...")
    fotos, total_fotos = fotos_por_especie(llave_de)

    # ---------- Excel ----------
    wb = openpyxl.Workbook()
    res = wb.active
    res.title = "Resumen"
    por_jardin = defaultdict(list)
    for a in arboles:
        por_jardin[a["jardin"]].append(a)
    num = lambda j: int(j.split(" - ")[0])  # noqa: E731
    jardines = sorted(por_jardin, key=num)
    negrita, fondo = Font(bold=True), PatternFill("solid", fgColor="E6E9DC")
    res.append(["Jardín", "Árboles", "Especies", "Especies con foto nuestra", "Árboles de especies con foto",
                "Árboles con nombre 2014 ≠ 2018", "Árboles sin especie"])
    for j in jardines:
        aa = por_jardin[j]
        esp = {a["especie"] for a in aa if a["especie"]}
        res.append([j, len(aa), len(esp), sum(1 for e in esp if llave_de(e) in fotos),
                    sum(1 for a in aa if a["especie"] and llave_de(a["especie"]) in fotos),
                    sum(1 for a in aa if a["aviso"]), sum(1 for a in aa if not a["especie"])])
    res.append([])
    res.append(["Casa ITESO Clavígero se deja para el final (fuera del campus). Jardín Contemplativo: ver su hoja, "
                "pendiente de identificación por Juan Gabriel."])

    esp_hoja = wb.create_sheet("Especies por jardín")
    esp_hoja.append(["Jardín", "Especie", "Nombres comunes en el inventario", "Árboles", "Placas",
                     "Fotos nuestras (confirmadas)", "Partes con foto", "Ficha en el sitio"])
    guia_jardines = []
    for j in jardines:
        aa = por_jardin[j]
        grupos = defaultdict(list)
        for a in aa:
            grupos[a["especie"] or f"(sin especie) {a['comun_2014']}"].append(a)
        especies = []
        for e, lista in sorted(grupos.items(), key=lambda t: (-len(t[1]), t[0])):
            k = llave_de(e)
            comunes = Counter(a["comun_2018"] or a["comun_2014"] for a in lista)
            partes = [p for p in PARTES if p in fotos.get(k, {})]
            esp_hoja.append([j, e, ", ".join(c for c, _ in comunes.most_common()), len(lista),
                             ", ".join(a["placa"] for a in lista), total_fotos.get(k, 0), ", ".join(partes),
                             fichas.get(k, {}).get("url", "")])
            especies.append({"especie": e, "llave": k, "comunes": [c for c, _ in comunes.most_common()],
                             "arboles": [{"placa": a["placa"], "aviso": a["aviso"]} for a in lista],
                             "n_fotos": total_fotos.get(k, 0)})
        guia_jardines.append({"nombre": j, "n": len(aa), "especies": especies})

        ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "", j)[:31])
        ws.append(["Placa", "Especie (resuelta)", "Nombre común 2018", "Nombre común 2014", "Nombre científico 2018",
                   "Aviso de nombre", "Año plantación", "Altura 2018 (m)", "Diámetro fuste 2018 (cm)",
                   "Observación 2018", "Latitud", "Longitud", "Verificación", "Especie real (si es otra)", "Notas"])
        for a in sorted(aa, key=lambda a: int(a["placa"]) if a["placa"].isdigit() else 0):
            ws.append([int(a["placa"]) if a["placa"].isdigit() else a["placa"], a["especie"], a["comun_2018"],
                       a["comun_2014"], a["cientifico_2018"], a["aviso"], a["anio"], a["altura"], a["diam"],
                       a["obs"], float(a["lat"]), float(a["lon"]), "", "", ""])
        dv = DataValidation(type="list", formula1='"Está,No está,Otra especie,No se encontró la placa"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"M2:M{ws.max_row}")

    # Por fotografiar: especies sin alguna parte básica (completo, tronco, hoja); primero las de Arboretum II,
    # que concentra la mayoría y está junto al Edificio V
    pf = wb.create_sheet("Por fotografiar", 1)
    pf.append(["Especie", "Nombre común", "Árboles en campus", "Falta completo", "Falta tronco", "Falta hoja",
               "Falta flor", "Falta fruto", "Jardines (árboles)", "Placas en Arboretum II", "Otras placas"])
    por_especie = defaultdict(list)
    for a in arboles:
        if a["especie"]:
            por_especie[a["especie"]].append(a)
    pendientes = []
    for e, lista in por_especie.items():
        tiene = fotos.get(llave_de(e), {})
        falta = [p not in tiene for p in PARTES]
        if any(falta[:3]):
            arb2 = [a["placa"] for a in lista if a["jardin"] == "7 - Arboretum II"]
            pendientes.append((not arb2, -sum(falta[:3]), -len(lista), e, lista, falta, arb2))
    for _, _, _, e, lista, falta, arb2 in sorted(pendientes):
        js = Counter(a["jardin"] for a in lista)
        otras = [a["placa"] for a in lista if a["jardin"] != "7 - Arboretum II"][:10]
        pf.append([e, Counter(a["comun_2018"] or a["comun_2014"] for a in lista).most_common(1)[0][0], len(lista),
                   *["sí" if f else "" for f in falta], ", ".join(f"{j} ({n})" for j, n in js.most_common()),
                   ", ".join(arb2), ", ".join(otras)])

    filas_c, i = leer_contemplativo()
    wc = wb.create_sheet("Jardín Contemplativo")
    wc.append(["Núm. foto SW Maps", "Archivo de foto", "Latitud", "Longitud", "Pl@ntNet 1", "Score 1",
               "Pl@ntNet 2", "Score 2", "Identificación de Juan Gabriel", "Ficha asignada", "Notas"])
    for f in filas_c:
        wc.append([f[i["num"]], f[i["archivo_foto"]], f[i["latitud"]], f[i["longitud"]], f[i["sugerida_1"]],
                   f[i["score_1"]], f[i["sugerida_2"]], f[i["score_2"]], "", "", ""])

    for ws in wb.worksheets:
        for c in ws[1]:
            c.font, c.fill, c.alignment = negrita, fondo, Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        if ws.max_row > 1 and ws.title != "Resumen":
            ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            largo = max(len(str(c.value or "")) for c in col[:200])
            ws.column_dimensions[col[0].column_letter].width = min(max(10, largo + 2), 45)
    wb.save(XLSX_SALIDA)

    # ---------- Guía HTML ----------
    llaves = {e["llave"] for j in guia_jardines for e in j["especies"]}
    datos = {
        "jardines": guia_jardines,
        "fotos": {k: v for k, v in fotos.items() if k in llaves},
        "fichas": {k: v for k, v in fichas.items() if k in llaves},
        "partes": PARTES,
    }
    html = PLANTILLA.read_text(encoding="utf-8").replace(
        "/*__DATOS__*/null", json.dumps(datos, ensure_ascii=False, separators=(",", ":")))
    HTML_SALIDA.write_text(html, encoding="utf-8")

    for j in guia_jardines:
        con_foto = sum(1 for e in j["especies"] if e["llave"] in fotos)
        print(f"  {j['nombre']}: {j['n']} árboles, {len(j['especies'])} especies ({con_foto} con foto)")
    print(f"Escrito {XLSX_SALIDA}")
    print(f"Escrito {HTML_SALIDA} ({HTML_SALIDA.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    construir()

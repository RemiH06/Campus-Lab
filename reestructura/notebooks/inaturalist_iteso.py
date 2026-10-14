# -*- coding: utf-8 -*-
"""Cruce de las especies de Campus Lab con las observaciones de ITESO en iNaturalist.

1. Baja TODAS las observaciones (cualquier grado) del proyecto Naturalista ITESO y de los
   predios del ITESO en el Bosque La Primavera, con sus fotos y licencias.
   Ojo: los árboles y plantas de jardín del campus son cultivados, e iNaturalist deja lo
   cultivado en grado "casual" aunque la identificación sea correcta. Por eso no se filtra
   por grado de investigación al bajar; se guarda el grado y los acuerdos de identificación.
2. Resuelve cada nombre nuestro (lista oficial, árboles, especies con foto confirmada) a su
   taxón vigente en iNaturalist (sinónimos: Acacia pennatula -> Vachellia pennatula).
   Usa primero el enlace de `referencia_nombres_comunes.pagina_naturalista` si existe.
3. Excel: especies de la lista oficial sin foto confirmada nuestra que sí tienen
   observaciones de ITESO en iNaturalist, y las que no tienen; más las especies de iNaturalist
   ITESO que no están en la lista oficial.

Solo usa la API de lectura (abierta, sin token). 1 petición por segundo. Resumible: la
resolución de nombres se salta lo ya resuelto; la descarga de observaciones se rehace
completa cada vez (son ~20 peticiones) para refrescar grados.

Tablas: inat_observaciones, inat_fotos, inat_nombres_resueltos.
Salida: reestructura/salida/inaturalist_iteso/Especies_iNaturalist_ITESO.xlsx

Uso: python inaturalist_iteso.py            (todo)
     python inaturalist_iteso.py --solo-excel
"""
import csv
import json
import re
import sqlite3
import sys
import time
import unicodedata
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
BASE = Path(__file__).resolve().parents[2]
RUTA_BD = BASE / "experimentos/bd/campuslab_volumen.db"
RUTA_LISTA_ARBOLES = BASE / "reestructura/notebooks/Lista_Arboles_ITESO.csv"
RUTA_FICHAS = BASE / "reestructura/salida/sitio_campuslab/fichas.json"
SALIDA = BASE / "reestructura/salida/inaturalist_iteso"

API = "https://api.inaturalist.org/v1"
FUENTES = {"proyecto": "project_id=146508", "predios_blp": "place_id=185911"}
RANGOS_ESPECIE = {"species", "hybrid", "subspecies", "variety", "form"}
LICENCIAS_ABIERTAS = {"cc0", "cc-by", "cc-by-sa"}


def binomio(nombre):
    plano = unicodedata.normalize("NFKD", nombre or "").encode("ascii", "ignore").decode()
    # "Hibiscus × rosa-sinensis": la marca de híbrido no es parte del binomio
    # con guion: "rosa-sinensis" cortado en "rosa" se confundía con otras especies
    palabras = [w for w in re.findall(r"[a-z]+(?:-[a-z]+)*", plano.lower().replace("×", " x ")) if w != "x"]
    return " ".join(palabras[:2]) if len(palabras) >= 2 else ""


def binomio_completo(nombre):
    plano = unicodedata.normalize("NFKD", nombre or "").encode("ascii", "ignore").decode()
    return [w for w in re.findall(r"[a-z]+(?:-[a-z]+)*", plano.lower()) if w not in ("x", "var", "subsp", "ssp", "f")]


def get(ruta):
    time.sleep(1)
    req = urllib.request.Request(API + ruta, headers={"User-Agent": "CampusLab-ITESO/1.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.load(r)


def crear_tablas(c):
    c.executescript("""
        CREATE TABLE IF NOT EXISTS inat_observaciones (
            obs_id INTEGER PRIMARY KEY, fuentes TEXT, taxon_id INTEGER, taxon_nombre TEXT,
            taxon_rango TEXT, reino_iconico TEXT, grado TEXT, cultivada INTEGER,
            comunidad_taxon_id INTEGER, acuerdos INTEGER, desacuerdos INTEGER,
            usuario_login TEXT, usuario_nombre TEXT, observada_en TEXT, lat REAL, lng REAL,
            lugar TEXT, url TEXT, fecha_descarga TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS inat_fotos (
            foto_inat_id INTEGER PRIMARY KEY, obs_id INTEGER, posicion INTEGER,
            url_square TEXT, licencia TEXT, atribucion TEXT, ancho INTEGER, alto INTEGER);
        CREATE TABLE IF NOT EXISTS inat_nombres_resueltos (
            binomio TEXT PRIMARY KEY, nombre_original TEXT, taxon_id INTEGER,
            nombre_inat TEXT, rango TEXT, metodo TEXT,
            fecha TEXT DEFAULT CURRENT_TIMESTAMP);
    """)
    c.commit()


def bajar_observaciones(c):
    for fuente, filtro in FUENTES.items():
        ultimo, n = 0, 0
        while True:
            d = get(f"/observations?{filtro}&order_by=id&order=asc&id_above={ultimo}&per_page=200")
            if not d["results"]:
                break
            for o in d["results"]:
                guardar_observacion(c, o, fuente)
                ultimo = o["id"]
                n += 1
            print(f"  {fuente}: {n} (faltan {d['total_results'] - len(d['results'])})", flush=True)
        print(f"{fuente}: {n} observaciones")


def guardar_observacion(c, o, fuente):
    t = o.get("taxon") or {}
    lat = lng = None
    if o.get("location"):
        lat, lng = (float(x) for x in o["location"].split(","))
    usuario = o.get("user") or {}
    c.execute("""
        INSERT INTO inat_observaciones (obs_id, fuentes, taxon_id, taxon_nombre, taxon_rango,
            reino_iconico, grado, cultivada, comunidad_taxon_id, acuerdos, desacuerdos,
            usuario_login, usuario_nombre, observada_en, lat, lng, lugar, url)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(obs_id) DO UPDATE SET
            fuentes = CASE WHEN instr(fuentes, excluded.fuentes) THEN fuentes
                           ELSE fuentes || ',' || excluded.fuentes END,
            taxon_id=excluded.taxon_id, taxon_nombre=excluded.taxon_nombre,
            taxon_rango=excluded.taxon_rango, reino_iconico=excluded.reino_iconico,
            grado=excluded.grado, cultivada=excluded.cultivada,
            comunidad_taxon_id=excluded.comunidad_taxon_id, acuerdos=excluded.acuerdos,
            desacuerdos=excluded.desacuerdos, fecha_descarga=CURRENT_TIMESTAMP
    """, (o["id"], fuente, t.get("id"), t.get("name"), t.get("rank"), t.get("iconic_taxon_name"),
          o.get("quality_grade"), int(bool(o.get("captive"))), o.get("community_taxon_id"),
          o.get("num_identification_agreements"), o.get("num_identification_disagreements"),
          usuario.get("login"), usuario.get("name"), o.get("observed_on"), lat, lng,
          o.get("place_guess"), o.get("uri")))
    for i, f in enumerate(o.get("photos") or []):
        dims = f.get("original_dimensions") or {}
        c.execute("""INSERT OR REPLACE INTO inat_fotos VALUES (?,?,?,?,?,?,?,?)""",
                  (f["id"], o["id"], i, f.get("url"), f.get("license_code"), f.get("attribution"),
                   dims.get("width"), dims.get("height")))
    c.commit()


def nuestros_nombres(c):
    """binomio -> (nombre original, grupo, nombre común, taxon_id de pagina_naturalista)."""
    nombres = {}
    for grupo, comun, cientifico, pagina in c.execute(
            "SELECT grupo, nombre_comun, nombre_cientifico, pagina_naturalista FROM referencia_nombres_comunes"):
        m = re.search(r"/taxa/(\d+)", pagina or "")
        nombres.setdefault(binomio(cientifico), (cientifico, grupo, comun, int(m.group(1)) if m else None))
    with open(RUTA_LISTA_ARBOLES, encoding="utf-8") as f:
        for fila in csv.DictReader(f):
            nombres.setdefault(binomio(fila["nombre_cientifico"]), (fila["nombre_cientifico"], "arboles", fila["nombre_comun"], None))
    if RUTA_FICHAS.exists():
        for ficha in json.loads(RUTA_FICHAS.read_text(encoding="utf-8")).values():
            if ficha.get("tipo") == "arbol":
                campos = ficha.get("campos", {})
                nombres.setdefault(binomio(campos.get("nombre_cientifico")),
                                   (campos.get("nombre_cientifico"), "arboles", campos.get("nombre_comun"), None))
    oficiales = set(nombres) - {""}
    for (cientifico,) in c.execute("""SELECT DISTINCT e.nombre_cientifico FROM fotos f
                                      JOIN especies e ON e.id=f.especie_id WHERE f.estado='usable'"""):
        nombres.setdefault(binomio(cientifico), (cientifico, None, None, None))
    nombres.pop("", None)
    return nombres, oficiales


def resolver_nombres(c, nombres):
    hechos = {r[0] for r in c.execute("SELECT binomio FROM inat_nombres_resueltos")}
    pendientes = [b for b in nombres if b not in hechos]
    print(f"Nombres por resolver: {len(pendientes)} (ya resueltos: {len(hechos)})")

    # 1) por el enlace de Naturalista de la lista oficial, en lotes
    con_id = [b for b in pendientes if nombres[b][3]]
    for i in range(0, len(con_id), 30):
        lote = con_id[i:i + 30]
        taxa = {t["id"]: t for t in get("/taxa/" + ",".join(str(nombres[b][3]) for b in lote))["results"]}
        for b in lote:
            t = taxa.get(nombres[b][3])
            if t and not t.get("is_active") and t.get("current_synonymous_taxon_ids"):
                t = get(f"/taxa/{t['current_synonymous_taxon_ids'][0]}")["results"][0]
            guardar_nombre(c, b, nombres[b][0], t, "enlace_naturalista" if t else "enlace_roto")
    # 2) por búsqueda de texto (resuelve sinónimos al taxón vigente)
    ya_con_taxon = {r[0] for r in c.execute("SELECT binomio FROM inat_nombres_resueltos WHERE taxon_id IS NOT NULL")}
    for n, b in enumerate(p for p in pendientes if p not in ya_con_taxon):
        # el binomio y no el nombre original: "Annona muricata L" con el autor no encuentra nada
        d = get("/taxa?" + urllib.parse.urlencode({"q": b, "per_page": 10}))
        candidatos = [t for t in d["results"] if t.get("rank") in RANGOS_ESPECIE
                      and binomio(t.get("matched_term")) == b]
        # primero el nombre tal cual (sin variedad: "H. rosa-sinensis var. schizopetalus" es otra
        # especie en iNat), luego especie o híbrido antes que rangos menores
        candidatos.sort(key=lambda t: (len(binomio_completo(t.get("matched_term"))) > 2,
                                       t["rank"] not in ("species", "hybrid")))
        t = candidatos[0] if candidatos else None
        guardar_nombre(c, b, nombres[b][0], t, ("sinonimo" if binomio(t["name"]) != b else "exacto") if t else "sin_match")
        if n % 50 == 0:
            print(f"  búsqueda: {n}", flush=True)


def guardar_nombre(c, b, original, t, metodo):
    c.execute("INSERT OR REPLACE INTO inat_nombres_resueltos (binomio, nombre_original, taxon_id, nombre_inat, rango, metodo) VALUES (?,?,?,?,?,?)",
              (b, original, t and t["id"], t and t["name"], t and t["rank"], metodo))
    c.commit()


def confianza(o):
    if o["grado"] == "research":
        return 3, "Investigación"
    if (o["acuerdos"] or 0) >= 2 and not o["desacuerdos"]:
        return 2, "Comunidad de acuerdo" + (" (cultivada)" if o["cultivada"] else "")
    if o["grado"] == "needs_id":
        return 1, "Necesita identificación"
    return 0, "Casual sin acuerdo" + (" (cultivada)" if o["cultivada"] else "")


def hoja(wb, titulo, encabezados, filas):
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for f in filas:
        ws.append(f)
    for i, enc in enumerate(encabezados, start=1):
        ancho = max([len(str(enc))] + [len(str(f[i - 1] or "")) for f in filas[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(ancho + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def excel(c, nombres, oficiales):
    c.row_factory = sqlite3.Row
    a_inat = {r["binomio"]: r for r in c.execute("SELECT * FROM inat_nombres_resueltos")}
    confirmadas = {binomio(n) for (n,) in c.execute("""SELECT e.nombre_cientifico FROM fotos f
                       JOIN especies e ON e.id=f.especie_id WHERE f.estado='usable'""")}
    # observaciones agrupadas por binomio del taxón de iNaturalist (variedades -> especie)
    obs = defaultdict(list)
    for o in c.execute("SELECT * FROM inat_observaciones WHERE taxon_rango IN ('species','hybrid','subspecies','variety','form')"):
        obs[binomio(o["taxon_nombre"])].append(o)
    fotos = defaultdict(list)
    for f in c.execute("SELECT obs_id, licencia FROM inat_fotos"):
        fotos[f["obs_id"]].append(f["licencia"])

    def resumen(lista):
        lista = sorted(lista, key=lambda o: -confianza(o)[0])
        grados = Counter(confianza(o)[1].split(" (")[0] for o in lista)
        lic = Counter("abierta" if l in LICENCIAS_ABIERTAS else ("no comercial" if l and "nc" in l else "reservada")
                      for o in lista for l in fotos[o["obs_id"]])
        return [confianza(lista[0])[1], len(lista), grados.get("Investigación", 0),
                grados.get("Comunidad de acuerdo", 0), grados.get("Necesita identificación", 0),
                lic.get("abierta", 0), lic.get("no comercial", 0), lic.get("reservada", 0),
                ", ".join(sorted({o["usuario_login"] for o in lista})[:5]),
                lista[0]["url"], ", ".join(sorted({f for o in lista for f in o["fuentes"].split(",")}))]

    cols_resumen = ["mejor_confianza", "observaciones", "investigacion", "comunidad_de_acuerdo",
                    "necesita_id", "fotos_licencia_abierta", "fotos_no_comercial",
                    "fotos_todos_derechos", "autores_inat", "mejor_observacion", "fuente"]
    orden_grupo = {"arboles": 0, "aves": 1, "orquideas": 2, "suculentas": 3, "flores_primavera": 4}

    con_inat, sin_inat = [], []
    for b in sorted(oficiales - confirmadas, key=lambda x: (orden_grupo.get(nombres[x][1], 9), x)):
        original, grupo, comun, _ = nombres[b]
        r = a_inat.get(b)
        clave = binomio(r["nombre_inat"]) if r and r["nombre_inat"] else b
        nombre_inat = r["nombre_inat"] if r else None
        base = [grupo, comun, original, nombre_inat if nombre_inat and binomio(nombre_inat) != b else "",
                r["metodo"] if r else "sin_resolver"]
        if obs.get(clave):
            con_inat.append(base + resumen(obs[clave]))
        else:
            sin_inat.append(base)

    claves_nuestras = {binomio(r["nombre_inat"]) for r in a_inat.values() if r["nombre_inat"]} | set(nombres)
    fuera = [[obs[k][0]["reino_iconico"], obs[k][0]["taxon_nombre"]] + resumen(obs[k])
             for k in sorted(obs) if k not in claves_nuestras]
    fuera.sort(key=lambda f: (f[0] or "", f[1]))

    wb = Workbook()
    leeme = wb.active
    leeme.title = "Léeme"
    for linea in [
        "Cruce de especies de Campus Lab con observaciones de ITESO en iNaturalist",
        f"Generado: {time.strftime('%d-%m-%Y')}. Fuentes: proyecto Naturalista ITESO (146508) y lugar 'Predios del ITESO en el Bosque La Primavera' (185911).",
        "",
        "Faltantes con iNat: especies de la lista oficial sin foto confirmada nuestra que sí tienen observaciones.",
        "Faltantes sin iNat: especies de la lista oficial sin foto nuestra ni observación de ITESO (lista por fotografiar).",
        "iNat fuera de lista: especies observadas en ITESO que no están en la lista oficial.",
        "",
        "Confianza: Investigación = grado de investigación de iNaturalist.",
        "Comunidad de acuerdo = 2+ identificaciones coinciden y ninguna en contra, pero quedó 'casual' (casi siempre por ser cultivada: árboles y plantas de jardín nunca llegan a investigación).",
        "Necesita identificación / Casual sin acuerdo = solo pista, no confirma la especie.",
        "",
        "Licencias de las fotos: abierta (cc0, cc-by, cc-by-sa) se puede publicar dando crédito; no comercial (cc-by-nc...) probablemente sí para un sitio universitario sin fines de lucro, confirmar; todos los derechos = pedir permiso al autor.",
        "nombre_inat solo aparece si iNaturalist usa otro nombre (sinónimo o taxonomía nueva).",
    ]:
        leeme.append([linea])
    leeme.column_dimensions["A"].width = 140
    base_cols = ["grupo", "nombre_comun", "nombre_lista", "nombre_inat", "resolucion"]
    hoja(wb, "Faltantes con iNat", base_cols + cols_resumen, con_inat)
    hoja(wb, "Faltantes sin iNat", base_cols, sin_inat)
    hoja(wb, "iNat fuera de lista", ["reino", "nombre_inat"] + cols_resumen, fuera)
    SALIDA.mkdir(parents=True, exist_ok=True)
    destino = SALIDA / "Especies_iNaturalist_ITESO.xlsx"
    wb.save(destino)
    print(f"Excel: {destino}")
    print(f"Oficiales: {len(oficiales)}, sin foto confirmada: {len(oficiales - confirmadas)}, "
          f"con iNat: {len(con_inat)}, sin iNat: {len(sin_inat)}, iNat fuera de lista: {len(fuera)}")
    print("Con iNat por grupo y confianza:", Counter((f[0], f[5]) for f in con_inat).most_common())


if __name__ == "__main__":
    c = sqlite3.connect(RUTA_BD)
    crear_tablas(c)
    nombres, oficiales = nuestros_nombres(c)
    if "--solo-excel" not in sys.argv:
        bajar_observaciones(c)
        resolver_nombres(c, nombres)
    excel(c, nombres, oficiales)

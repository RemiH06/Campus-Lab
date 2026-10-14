# -*- coding: utf-8 -*-
"""Paso 3 del cruce con iNaturalist ITESO: buscar la MISMA foto en nuestro catálogo y en
iNaturalist, y comparar la especie que dice cada lado.

No busca fotos "parecidas a la especie": solo pares casi idénticos (criterio de Remi, 13-oct),
porque la idea es usar la identificación de iNaturalist como control de las nuestras
(¿clasificamos o confirmamos mal?).

Un par cuenta si:
  - similitud BioCLIP >= 0.95 (mismo modelo y preprocesado que Clustering_Especimen.ipynb), o
  - dHash a <= 4 bits de 64 y similitud >= 0.60 (misma imagen rotada: nuestros embeddings se
    calcularon sin enderezar por EXIF y las de iNaturalist vienen derechas), o
  - dHash a <= 10 bits y similitud >= 0.85.
  La huella sola no basta: a 8-10 bits empareja cosas sin relación (un chapulín con una zorra).

Entradas: especimen_embeddings (nuestras), tablas inat_* (inaturalist_iteso.py) y
salida/inaturalist_iteso/pixilate/embeddings_inat.npz (embeddings_inat.py, corre en Pixilate).
Salida: salida/inaturalist_iteso/Cruce_fotos_iNaturalist.xlsx

Uso: python cruce_fotos_inat.py --hashes   (dHash de nuestras fotos, cache resumible)
     python cruce_fotos_inat.py            (cruce y Excel)
"""
import json
import re
import sqlite3
import sys
import unicodedata
from collections import Counter
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

sys.stdout.reconfigure(encoding="utf-8")
BASE = Path(__file__).resolve().parents[2]
RUTA_BD = BASE / "experimentos/bd/campuslab_volumen.db"
RUTA_CONFIG = BASE / "reestructura/notebooks/config_local.json"
CARPETA = BASE / "reestructura/salida/inaturalist_iteso"
CACHE_HASH = CARPETA / "dhash_nuestras.json"
NPZ_INAT = CARPETA / "pixilate/embeddings_inat.npz"
UMBRAL_SIM = 0.95
UMBRAL_HASH = 10
UMBRAL_HASH_ESTRICTO = 4


def binomio(nombre):
    plano = unicodedata.normalize("NFKD", nombre or "").encode("ascii", "ignore").decode()
    # "Hibiscus × rosa-sinensis": la marca de híbrido no es parte del binomio
    # con guion: "rosa-sinensis" cortado en "rosa" se confundía con otras especies
    palabras = [w for w in re.findall(r"[a-z]+(?:-[a-z]+)*", plano.lower().replace("×", " x ")) if w != "x"]
    return " ".join(palabras[:2]) if len(palabras) >= 2 else ""


def dhash(img, tam=8):
    """Igual que en embeddings_inat.py (Pixilate): si cambia uno, cambiar el otro."""
    gris = ImageOps.exif_transpose(img).convert("L").resize((tam + 1, tam), Image.LANCZOS)
    p = np.asarray(gris, dtype=np.int16)
    bits = (p[:, 1:] > p[:, :-1]).flatten()
    return int("".join("1" if b else "0" for b in bits), 2)


def calcular_hashes(c):
    espejo = Path(json.loads(RUTA_CONFIG.read_text(encoding="utf-8"))["ruta_espejo_local_fotos"])
    cache = json.loads(CACHE_HASH.read_text(encoding="utf-8")) if CACHE_HASH.exists() else {}
    filas = c.execute("""SELECT e.foto_id, f.ruta_actual FROM especimen_embeddings e
                         JOIN fotos f ON f.id = e.foto_id""").fetchall()
    pendientes = [(i, r) for i, r in filas if str(i) not in cache]
    print(f"dHash: {len(filas)} fotos, faltan {len(pendientes)}")
    for n, (foto_id, ruta) in enumerate(pendientes, 1):
        origen = espejo / Path(ruta).name
        if not origen.exists():
            origen = Path(ruta)
        try:
            with Image.open(origen) as img:
                img.draft("RGB", (600, 600))  # JPEG: decodifica reducido, mucho más rápido
                cache[str(foto_id)] = str(dhash(img))
        except Exception as e:
            print(f"  error {foto_id}: {e}")
        if n % 250 == 0 or n == len(pendientes):
            CACHE_HASH.write_text(json.dumps(cache), encoding="utf-8")
            print(f"  {n}/{len(pendientes)}", flush=True)


def distancias_hash(a, b):
    """Matriz de distancias de Hamming entre dos listas de hashes de 64 bits."""
    bits_por_byte = np.array([bin(i).count("1") for i in range(256)], dtype=np.uint8)
    a, b = np.array(a, dtype=np.uint64), np.array(b, dtype=np.uint64)
    salida = np.empty((len(a), len(b)), dtype=np.uint8)
    for i in range(0, len(a), 500):  # por bloques: la matriz completa en bits no cabe en memoria
        x = a[i:i + 500, None] ^ b[None, :]
        salida[i:i + 500] = bits_por_byte[x.view(np.uint8).reshape(x.shape + (8,))].sum(axis=-1, dtype=np.uint8)
    return salida


def nuestra_identidad(c):
    """foto_id -> (nombre, estado_legible)."""
    info = {}
    for foto_id, estado, nombre in c.execute("""SELECT f.id, f.estado, e.nombre_cientifico
                                                FROM fotos f LEFT JOIN especies e ON e.id=f.especie_id"""):
        if nombre:
            info[foto_id] = (nombre, "confirmada" if estado == "usable" else f"asignada ({estado})")
    for foto_id, sugerida in c.execute("""SELECT foto_id, especie_sugerida FROM identificaciones_revisadas
                                          WHERE decision='rechazada' AND especie_sugerida IS NOT NULL
                                          ORDER BY fecha"""):
        if foto_id not in info:
            info[foto_id] = (sugerida, "limbo")
    return info


def cruce(c):
    c.row_factory = None
    hashes = json.loads(CACHE_HASH.read_text(encoding="utf-8"))
    nuestras = [(i, np.frombuffer(e, dtype=np.float32)) for i, e in
                c.execute("SELECT foto_id, embedding FROM especimen_embeddings")]
    ids_n = np.array([i for i, _ in nuestras])
    emb_n = np.stack([e for _, e in nuestras])
    hash_n = [int(hashes.get(str(i), 0)) for i in ids_n]

    z = np.load(NPZ_INAT)
    ids_i, emb_i, hash_i = z["ids"], z["emb"], [int(h) for h in z["dhash"]]
    sim = emb_n @ emb_i.T
    dist = distancias_hash(hash_n, hash_i)
    sin_hash = np.array([str(i) not in hashes for i in ids_n])
    dist[sin_hash, :] = 64
    pares = np.argwhere((sim >= UMBRAL_SIM) | ((dist <= UMBRAL_HASH_ESTRICTO) & (sim >= 0.60))
                        | ((dist <= UMBRAL_HASH) & (sim >= 0.85)))
    print(f"Pares: {len(pares)} (BioCLIP>={UMBRAL_SIM}: {(sim >= UMBRAL_SIM).sum()})")

    ident = nuestra_identidad(c)
    a_inat = {b: binomio(n) for b, n in c.execute(
        "SELECT binomio, nombre_inat FROM inat_nombres_resueltos WHERE nombre_inat IS NOT NULL")}
    obs_de_foto = {f: o for f, o in c.execute("SELECT foto_inat_id, obs_id FROM inat_fotos")}
    licencia = {f: l for f, l in c.execute("SELECT foto_inat_id, licencia FROM inat_fotos")}
    obs = {r[0]: r for r in c.execute("""SELECT obs_id, taxon_nombre, taxon_rango, grado, acuerdos,
                                                desacuerdos, cultivada, usuario_login, url, reino_iconico
                                         FROM inat_observaciones""")}
    grupo = dict(c.execute("SELECT foto_id, grupo_id FROM especimen_grupos_confirmados"))
    reino_genero = {g.lower(): r for g, r in c.execute("SELECT genero, reino FROM genero_taxonomia") if g}
    reino_iconico = {"Plantae": "Plantae", "Fungi": "Fungi"}  # lo demás (Aves, Insecta...) es Animalia

    filas = []
    for a, b in pares:
        foto_id, foto_inat = int(ids_n[a]), int(ids_i[b])
        o = obs[obs_de_foto[foto_inat]]
        nombre, estado = ident.get(foto_id, ("", "sin especie"))
        nuestro = binomio(nombre)
        nuestro = a_inat.get(nuestro, nuestro)  # llevar sinónimos al nombre vigente de iNat
        if o[2] in ("species", "hybrid", "subspecies", "variety", "form"):
            suyo = binomio(o[1])
            coincide = nuestro == suyo if nuestro else None
        else:  # iNat solo llega a género o más arriba: comparar el género
            suyo = (o[1] or "").lower()
            coincide = (nuestro.split(" ")[0] == suyo) if nuestro and o[2] == "genus" else None
        if not nuestro:
            categoria = "4. nuestra sin especie (iNat la sugiere)"
        elif coincide is None:
            categoria = "5. iNat no llega a especie/género"
        elif coincide:
            categoria = "1. coincide" + ("" if o[2] != "genus" else " (solo género)")
        elif (reino_genero.get(nuestro.split(" ")[0]) and o[9]
              and reino_genero[nuestro.split(" ")[0]] != reino_iconico.get(o[9], "Animalia")):
            categoria = "6. otro sujeto en la misma foto (reino distinto)"
        elif estado == "confirmada":
            categoria = "2. NO coincide, nuestra confirmada"
        else:
            categoria = "3. NO coincide, nuestra en limbo/asignada"
        filas.append([categoria, round(float(sim[a, b]), 3), int(dist[a, b]),
                      "sí" if dist[a, b] <= UMBRAL_HASH_ESTRICTO else "", foto_id, grupo.get(foto_id),
                      nombre, estado, o[1], o[2], o[3], o[4], o[5], "sí" if o[6] else "",
                      o[7], licencia.get(foto_inat), o[8],
                      f"http://localhost:8044/?id={foto_id}"])
    filas.sort(key=lambda f: (f[0], -f[1]))
    return filas


def excel(filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Léeme"
    for linea in [
        "Cruce foto a foto: nuestro catálogo vs. observaciones de ITESO en iNaturalist",
        f"Un par = similitud BioCLIP >= {UMBRAL_SIM}, o huella dHash a <= {UMBRAL_HASH_ESTRICTO}/64 bits con similitud >= 0.60 (foto rotada), o a <= {UMBRAL_HASH} bits con similitud >= 0.85. misma_imagen = sí si la huella está a <= {UMBRAL_HASH_ESTRICTO} bits.",
        "Similitud alta sin huella igual suele ser otra toma del mismo ejemplar (ráfaga); la identificación sigue sirviendo de control.",
        "Categoría 6: la foto tiene dos sujetos (planta y animal); nuestra fila es uno y la observación de iNat el otro. No es error: candidata a fila de doble sujeto.",
        "Categorías: 1 coincide · 2 NO coincide y nuestra está confirmada (revisar primero) · 3 NO coincide y la nuestra está en limbo o sin confirmar · 4 nosotros no tenemos especie · 5 iNat no llega a especie ni género.",
        "Los sinónimos se comparan con el nombre vigente de iNaturalist (Acacia pennatula = Vachellia pennatula).",
        "Confianza de iNat: grado research > acuerdos altos sin desacuerdos > needs_id/casual. Cultivada = sí explica por qué no llega a research.",
        "La columna revisar abre la foto en buscar_foto (8044, tiene que estar corriendo).",
    ]:
        ws.append([linea])
    ws.column_dimensions["A"].width = 140
    resumen = Counter(f[0] for f in filas)
    ws.append([])
    for k in sorted(resumen):
        ws.append([f"{k}: {resumen[k]} pares, {len({f[4] for f in filas if f[0] == k})} fotos nuestras"])

    enc = ["categoria", "similitud", "dist_hash", "misma_imagen", "foto_id", "grupo_id",
           "nuestra_especie", "nuestro_estado", "inat_taxon", "inat_rango", "inat_grado",
           "inat_acuerdos", "inat_desacuerdos", "inat_cultivada", "inat_autor", "licencia_foto",
           "observacion", "revisar"]
    ws = wb.create_sheet("Pares")
    ws.append(enc)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for f in filas:
        ws.append(f)
    for i, e in enumerate(enc, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max([len(e)] + [len(str(f[i - 1] or "")) for f in filas[:300]]) + 2, 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    destino = CARPETA / "Cruce_fotos_iNaturalist.xlsx"
    wb.save(destino)
    print(f"Excel: {destino}")
    for k in sorted(resumen):
        print(f"  {k}: {resumen[k]} pares, {len({f[4] for f in filas if f[0] == k})} fotos nuestras")


if __name__ == "__main__":
    con = sqlite3.connect(RUTA_BD)
    if "--hashes" in sys.argv:
        calcular_hashes(con)
    else:
        excel(cruce(con))
